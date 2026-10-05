"""事務所 LINE 官方帳號追蹤（mig 241）。

模式：
  python line_oa_daily.py                # discover + daily（排程用）
  python line_oa_daily.py discover       # 從 firm_digital_signals 官網 LINE 連結找新帳號
  python line_oa_daily.py daily          # 抓所有帳號主頁 → 更新認證狀態、寫當日好友數
  python line_oa_daily.py add <id> [所名] [web_search|manual]   # 手動加帳號
  python line_oa_daily.py seed-xlsx <競品追蹤表.xlsx>             # 一次性匯入帳號＋歷史

資料來源＝page.line.me/<id> 的 __NEXT_DATA__（badgeType / friendCount）。
只有認證帳號有公開主頁；未認證帳號回空殼頁（記 no_page、不寫好友數）。
friendCount 是淨好友數，日差＝新增−封鎖。
"""
import json
import os
import re
import sys
import urllib.parse
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta, timezone

import requests

HERE = os.path.dirname(os.path.abspath(__file__))


def load_env():
    env = dict(os.environ)
    p = os.path.join(HERE, ".env")
    if os.path.exists(p):
        with open(p, encoding="utf-8-sig") as f:
            for line in f:
                line = line.strip()
                if line and not line.startswith("#") and "=" in line:
                    k, v = line.split("=", 1)
                    env.setdefault(k.strip(), v.strip())
    return env


ENV = load_env()
SB_URL = ENV["SUPABASE_URL"].strip().rstrip("/")
SB_KEY = ENV["SUPABASE_SERVICE_KEY"].strip()
HDR = {"apikey": SB_KEY, "Authorization": "Bearer " + SB_KEY, "Content-Type": "application/json"}

UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/126.0.0.0 Safari/537.36")
LINE = requests.Session()
LINE.headers["User-Agent"] = UA
LINE.headers["Accept-Language"] = "zh-TW,zh;q=0.9"

TPE = timezone(timedelta(hours=8))
LAW_WORD = re.compile(r"法律|律師|法務|法|law|legal", re.I)


def today_tpe():
    return datetime.now(TPE).date().isoformat()


# ---------------- Supabase (PostgREST) ----------------
def sb_get(path, page=1000):
    out, start = [], 0
    while True:
        r = requests.get(f"{SB_URL}/rest/v1/{path}", timeout=60,
                         headers={**HDR, "Range-Unit": "items", "Range": f"{start}-{start + page - 1}"})
        r.raise_for_status()
        rows = r.json()
        out += rows
        if len(rows) < page:
            return out
        start += page


def sb_upsert(table, rows, on_conflict):
    for i in range(0, len(rows), 200):
        r = requests.post(f"{SB_URL}/rest/v1/{table}?on_conflict={on_conflict}", timeout=60,
                          headers={**HDR, "Prefer": "resolution=merge-duplicates,return=minimal"},
                          data=json.dumps(rows[i:i + 200], ensure_ascii=False).encode("utf-8"))
        if r.status_code >= 300:
            raise RuntimeError(f"upsert {table} HTTP {r.status_code}: {r.text[:300]}")


# ---------------- LINE ----------------
def norm_id(s):
    return (s or "").strip().lstrip("@").lower() or None


def resolve_id(url):
    """官網上的 LINE 連結 → 官方帳號 ID；個人帳號／非帳號連結回 None。"""
    u = urllib.parse.unquote(url or "")
    if "lin.ee/" in u:
        try:
            r = LINE.get(u, allow_redirects=False, timeout=15)
            u = urllib.parse.unquote(r.headers.get("Location", ""))
        except requests.RequestException:
            return None
    m = (re.search(r"line\.me/(?:R/)?ti/p/~?@([\w.\-]+)", u)
         or re.search(r"page\.line\.me/@?([\w.\-]+)", u))
    return m.group(1).lower() if m else None


def fetch_profile(search_id):
    """回 dict：page=False（無公開主頁＝未認證）或 page=True 加各欄位。網路錯誤丟例外。"""
    r = LINE.get(f"https://page.line.me/{search_id}", timeout=20)
    if r.status_code >= 500:
        raise RuntimeError(f"HTTP {r.status_code}")
    m = re.search(r'<script id="__NEXT_DATA__"[^>]*>(.*?)</script>', r.text, re.S)
    if not m:
        return {"page": False}
    d = json.loads(json.loads(m.group(1))["props"]["pageProps"]["initialDataString"])
    acc = d.get("account") or {}
    ai, pf = acc.get("accountInfo") or {}, acc.get("profile") or {}
    if not ai:
        return {"page": False}
    return {
        "page": True,
        "badge": pf.get("badgeType"),
        "name": pf.get("name"),
        "friends": ai.get("friendCount"),
        "basic_id": (ai.get("basicSearchId") or "").lower() or None,
        "premium_id": (ai.get("premiumSearchId") or "").lower() or None,
        "country": ai.get("countryCode"),
    }


# ---------------- modes ----------------
def load_firm_names():
    return {r["firm_name"] for r in sb_get("moj_firm_stats_cache?select=firm_name") if r.get("firm_name")}


def pick_firm(display_name, fallback, firm_names):
    """帳號顯示名含某名冊所名（取最長）→ 用它；否則沿用來源所名。"""
    hits = [f for f in firm_names if len(f) >= 5 and f in (display_name or "")]
    return max(hits, key=len) if hits else fallback


def discover():
    existing = sb_get("line_oa_accounts?select=search_id,basic_id")
    known = {r["search_id"] for r in existing} | {norm_id(r["basic_id"]) for r in existing if r["basic_id"]}
    sig = sb_get("firm_digital_signals?select=firm_name,line_url&line_url=not.is.null")
    cand = {}
    with ThreadPoolExecutor(6) as ex:
        for row, sid in zip(sig, ex.map(lambda x: resolve_id(x["line_url"]), sig)):
            if sid and sid not in known and sid not in cand:
                cand[sid] = row
    if not cand:
        print(f"discover: 官網 LINE 連結 {len(sig)} 筆，無新帳號")
        return
    firm_names = load_firm_names()
    new = []
    with ThreadPoolExecutor(6) as ex:
        profs = list(ex.map(lambda s: _safe_profile(s), cand))
    for sid, prof in zip(cand, profs):
        row = cand[sid]
        if prof.get("basic_id") and norm_id(prof["basic_id"]) in known:
            continue  # 同帳號另一個 ID 已收錄
        name = prof.get("name")
        rec = {"search_id": sid, "source": "website", "source_url": row["line_url"],
               "firm_name": pick_firm(name, row["firm_name"], firm_names) if name else row["firm_name"],
               "display_name": name, "basic_id": prof.get("basic_id"), "premium_id": prof.get("premium_id"),
               "is_certified": prof.get("badge") == "certified" if "page" in prof else None,
               "excluded": False, "note": None}
        if name and not LAW_WORD.search(name):
            rec["excluded"], rec["note"] = True, "顯示名不像法律帳號，待人工確認"
        known.add(sid)
        if prof.get("basic_id"):
            known.add(norm_id(prof["basic_id"]))
        new.append(rec)
    sb_upsert("line_oa_accounts", new, "search_id")
    print(f"discover: 新增 {len(new)} 個帳號（認證 {sum(1 for r in new if r['is_certified'])}，"
          f"排除 {sum(1 for r in new if r['excluded'])}）")


def _safe_profile(sid):
    try:
        return fetch_profile(sid)
    except Exception as e:  # noqa: BLE001
        return {"error": type(e).__name__}


def daily():
    accs = sb_get("line_oa_accounts?select=search_id,display_name,firm_name&excluded=eq.false")
    snap = today_tpe()
    now = datetime.now(timezone.utc).isoformat()
    with ThreadPoolExecutor(6) as ex:
        profs = list(ex.map(lambda a: _safe_profile(a["search_id"]), accs))
    meta, daily_rows, n_err = [], [], 0
    for a, p in zip(accs, profs):
        m = {"search_id": a["search_id"], "last_checked_at": now}
        if "error" in p:
            n_err += 1
            m["last_status"] = "error:" + p["error"]
        elif not p["page"]:
            m.update(last_status="no_page", is_certified=False)
        else:
            m.update(last_status="ok", is_certified=p["badge"] == "certified",
                     display_name=p["name"], basic_id=p["basic_id"], premium_id=p["premium_id"])
            if isinstance(p["friends"], int):
                daily_rows.append({"search_id": a["search_id"], "snap_date": snap,
                                   "friend_count": p["friends"], "source": "scraper", "fetched_at": now})
        meta.append(m)
    # PostgREST bulk upsert 要求每列欄位一致 → 依欄位組分批
    groups = {}
    for m in meta:
        groups.setdefault(tuple(sorted(m)), []).append(m)
    for rows in groups.values():
        sb_upsert("line_oa_accounts", rows, "search_id")
    sb_upsert("line_oa_daily", daily_rows, "search_id,snap_date")
    n_np = sum(1 for m in meta if m.get("last_status") == "no_page")
    print(f"daily {snap}: 帳號 {len(accs)}｜寫入好友數 {len(daily_rows)}｜無主頁(未認證) {n_np}｜錯誤 {n_err}")
    summary = os.environ.get("GITHUB_STEP_SUMMARY")
    if summary:
        with open(summary, "a", encoding="utf-8") as f:
            f.write(f"## LINE 官方帳號 {snap}\n\n- 帳號 {len(accs)}\n- 寫入好友數 {len(daily_rows)}\n"
                    f"- 無主頁（未認證）{n_np}\n- 錯誤 {n_err}\n")
    # 被擋防呆：錯誤過半或一筆好友數都沒寫到 → 紅燈
    if accs and (n_err * 2 > len(accs) or not daily_rows):
        print("::error::多數帳號抓取失敗，疑似被 LINE 擋或頁面改版")
        sys.exit(1)


def add(sid, firm=None, source="manual"):
    sid = norm_id(sid)
    p = fetch_profile(sid)
    rec = {"search_id": sid, "firm_name": firm, "source": source,
           "source_url": f"https://page.line.me/{sid}", "excluded": False,
           "is_certified": p.get("badge") == "certified" if p["page"] else False,
           "display_name": p.get("name"), "basic_id": p.get("basic_id"), "premium_id": p.get("premium_id")}
    sb_upsert("line_oa_accounts", [rec], "search_id")
    print("add:", sid, rec["display_name"], "認證" if rec["is_certified"] else "未認證/無主頁", p.get("friends"))


# 競品追蹤表的帳號名 → 名冊所名（顯示名對不到所名、需人工指定者）
SHEET_FIRM_OVERRIDE = {"第一法律": "第一國際法律事務所", "法律我幫您・雍和法律事務所": "雍和法律事務所"}


def seed_xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    firm_names = load_firm_names()
    accs, hist = {}, []
    for r in wb["最新狀態總覽"].iter_rows(min_row=4, values_only=True):
        _, cat, name, lid, _url = r[:5]
        sid = norm_id(lid)
        if not sid or sid in accs:
            continue
        accs[sid] = {"cat": cat, "name": name}
    for r in wb["歷史數據記錄"].iter_rows(min_row=3, values_only=True):
        d, cat, lid, name, cnt = r[:5]
        sid = norm_id(lid)
        if sid and isinstance(cnt, (int, float)) and d:
            hist.append((sid, d.date().isoformat(), int(cnt)))
            accs.setdefault(sid, {"cat": cat, "name": name})
    # 用主頁的 basic_id 合併同帳號不同 ID（例：聖安 @kxe6497n＝@twlaw）
    with ThreadPoolExecutor(6) as ex:
        profs = dict(zip(accs, ex.map(_safe_profile, accs)))
    canon, by_basic = {}, {}
    for sid, p in profs.items():
        b = norm_id(p.get("basic_id"))
        canon[sid] = by_basic.setdefault(b, sid) if b else sid
    existing = {r["search_id"]: r for r in sb_get("line_oa_accounts?select=search_id,basic_id")}
    ex_basic = {norm_id(r["basic_id"]): s for s, r in existing.items() if r["basic_id"]}
    rows = []
    for sid, info in accs.items():
        if canon[sid] != sid:
            continue
        p = profs[sid]
        b = norm_id(p.get("basic_id"))
        if b in ex_basic and ex_basic[b] != sid:
            canon[sid] = ex_basic[b]  # 官網 discover 已收錄同帳號 → 歷史併過去，不另建
            continue
        name = p.get("name") or info["name"]
        zl = info["cat"] == "喆律" or "喆律" in (name or "")
        firm = "喆律法律事務所" if zl else (SHEET_FIRM_OVERRIDE.get(name) or pick_firm(name, None, firm_names))
        rows.append({"search_id": sid, "display_name": name, "firm_name": firm,
                     "brand_group": "喆律" if zl else None, "category": None if zl else info["cat"],
                     "source": "tracker_sheet", "source_url": f"https://page.line.me/{sid}",
                     "is_certified": (p.get("badge") == "certified") if "page" in p else None,
                     "basic_id": p.get("basic_id"), "premium_id": p.get("premium_id"), "excluded": False})
    sb_upsert("line_oa_accounts", rows, "search_id")
    seen, hrows = set(), []
    for sid, d, cnt in hist:
        key = (canon.get(sid, sid), d)
        if key in seen:
            continue
        seen.add(key)
        hrows.append({"search_id": key[0], "snap_date": d, "friend_count": cnt, "source": "tracker_sheet"})
    sb_upsert("line_oa_daily", hrows, "search_id,snap_date")
    print(f"seed-xlsx: 帳號 {len(rows)}（喆律 {sum(1 for r in rows if r['brand_group'])}，"
          f"對到所名 {sum(1 for r in rows if r['firm_name'])}）｜歷史 {len(hrows)} 筆")


if __name__ == "__main__":
    mode = sys.argv[1] if len(sys.argv) > 1 else "all"
    if mode == "discover":
        discover()
    elif mode == "daily":
        daily()
    elif mode == "add":
        add(sys.argv[2], sys.argv[3] if len(sys.argv) > 3 else None, sys.argv[4] if len(sys.argv) > 4 else "manual")
    elif mode == "seed-xlsx":
        seed_xlsx(sys.argv[2])
    else:
        discover()
        daily()
